"""
Еженедельное планирование загрузки команды.
Запуск: streamlit run app.py
"""
import io
import re
from datetime import datetime

import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Загрузка команды", layout="wide")

OVERVIEW_PREFIX = "Обзор тикетов"
SUMMARY_SHEET_NAME = "Сводка по команде"
FIXED_COLS = ["Тикет", "Название", "Колонка", "Доска", "Срок", "В работе"]
DEFAULT_NORM = 40.0


# ---------------------------------------------------------------------------
# Parsing the weekly export from the Agile board
# ---------------------------------------------------------------------------

@st.cache_data(show_spinner=False)
def parse_workbook(file_bytes: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    person_sheets = [
        s for s in wb.sheetnames
        if not s.startswith(OVERVIEW_PREFIX) and s != SUMMARY_SHEET_NAME
    ]

    people_data = {}
    week_labels = None

    for name in person_sheets:
        ws = wb[name]
        header = [c.value for c in ws[2]]
        if "Тикет" not in header:
            continue
        idx = {h: i for i, h in enumerate(header) if h is not None}
        start_week_col = idx["В работе"] + 1
        end_week_col = idx.get("Итого, ч.", len(header) - 1)
        this_week_labels = [header[i] for i in range(start_week_col, end_week_col)]
        if week_labels is None:
            week_labels = this_week_labels

        rows = []
        for r in range(3, ws.max_row + 1):
            a_val = ws.cell(row=r, column=1).value
            c_val = ws.cell(row=r, column=idx["Колонка"] + 1).value
            if a_val == "Итого по неделям:":
                break
            row = {
                "Тикет": a_val or "",
                "Название": ws.cell(row=r, column=idx["Название"] + 1).value or "",
                "Колонка": c_val or "",
                "Доска": ws.cell(row=r, column=idx.get("Доска", 2) + 1).value if "Доска" in idx else "",
                "Срок": ws.cell(row=r, column=idx["Срок"] + 1).value,
                "В работе": bool(ws.cell(row=r, column=idx["В работе"] + 1).value),
            }
            for wi, wl in enumerate(this_week_labels):
                v = ws.cell(row=r, column=start_week_col + 1 + wi).value
                row[wl] = float(v) if isinstance(v, (int, float)) else 0.0
            rows.append(row)

        df = pd.DataFrame(rows, columns=FIXED_COLS + this_week_labels)
        people_data[name] = df

    return people_data, (week_labels or [])


def clean_week_label(label: str) -> str:
    return re.sub(r"\s*ч\.?$", "", str(label)).strip()


# ---------------------------------------------------------------------------
# Excel export (mirrors the manually-built workbook from earlier iterations)
# ---------------------------------------------------------------------------

def build_export_workbook(people_data: dict, week_labels: list, norms: dict) -> bytes:
    HEADER_FILL = PatternFill("solid", fgColor="2F4A6D")
    HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
    INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
    INPUT_FONT = Font(name="Arial", color="0000FF")
    LINK_FONT = Font(name="Arial", color="008000")
    TOTAL_FILL = PatternFill("solid", fgColor="DDE5EE")
    TOTAL_FONT = Font(name="Arial", bold=True)
    BODY_FONT = Font(name="Arial")
    TITLE_FONT = Font(name="Arial", bold=True, size=14)
    THIN = Side(style="thin", color="B7B7B7")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    OVERLOAD_FILL = PatternFill(bgColor="F4CCCC")
    OK_FILL = PatternFill(bgColor="D9EAD3")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    totals_row = {}
    for name, df in people_data.items():
        ws = wb.create_sheet(name[:31])
        ws.cell(row=1, column=1, value=f"{name} — часы по неделям, заполнено в планировщике")
        headers = FIXED_COLS + week_labels + ["Итого, ч."]
        for ci, h in enumerate(headers, start=1):
            c = ws.cell(row=2, column=ci, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

        r = 3
        for _, row in df.iterrows():
            for ci, col in enumerate(FIXED_COLS, start=1):
                ws.cell(row=r, column=ci, value=row[col])
            week_start_col = len(FIXED_COLS) + 1
            for wi, wl in enumerate(week_labels):
                ws.cell(row=r, column=week_start_col + wi, value=row[wl])
            total_col = week_start_col + len(week_labels)
            ws.cell(row=r, column=total_col,
                     value=f"=SUM({get_column_letter(week_start_col)}{r}:{get_column_letter(total_col - 1)}{r})")
            r += 1

        last_data_row = r - 1
        ws.cell(row=r, column=1, value="Итого по неделям:").font = TOTAL_FONT
        week_start_col = len(FIXED_COLS) + 1
        total_col = week_start_col + len(week_labels)
        for ci in range(week_start_col, total_col + 1):
            letter = get_column_letter(ci)
            c = ws.cell(row=r, column=ci, value=f"=SUM({letter}3:{letter}{last_data_row})")
            c.font = TOTAL_FONT
            c.fill = TOTAL_FILL
        ws.cell(row=r, column=1).fill = TOTAL_FILL
        totals_row[name] = r

        widths = [14, 50, 22, 10, 12, 10] + [14] * len(week_labels) + [12]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --- summary sheet ---
    ws = wb.create_sheet(SUMMARY_SHEET_NAME, 0)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "Еженедельное планирование загрузки команды"
    ws["A1"].font = TITLE_FONT

    header_row = 3
    headers = ["Исполнитель", "Норма, ч/нед"] + week_labels + ["Итого, ч", "Загрузка посл. недели"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    first_data_row = header_row + 1
    names = list(people_data.keys())
    for idx, name in enumerate(names):
        r = first_data_row + idx
        ws.cell(row=r, column=1, value=name).font = BODY_FONT
        norm_cell = ws.cell(row=r, column=2, value=norms.get(name, DEFAULT_NORM))
        norm_cell.fill = INPUT_FILL
        norm_cell.font = INPUT_FONT
        tr = totals_row[name]
        week_start_col = len(FIXED_COLS) + 1
        for wi in range(len(week_labels)):
            cell = ws.cell(row=r, column=3 + wi)
            cell.value = f"='{name}'!{get_column_letter(week_start_col + wi)}{tr}"
            cell.font = LINK_FONT
            cell.number_format = "0.#"
        last_week_col = 2 + len(week_labels)
        total_col = last_week_col + 1
        ws.cell(row=r, column=total_col, value=f"=SUM(C{r}:{get_column_letter(last_week_col)}{r})").font = BODY_FONT
        load_col = total_col + 1
        ws.cell(row=r, column=load_col,
                value=f'=IF(B{r}=0,"",{get_column_letter(last_week_col)}{r}/B{r})').number_format = "0%"

    last_data_row = first_data_row + len(names) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="Итого по команде:").font = TOTAL_FONT
    last_week_col = 2 + len(week_labels)
    total_col = last_week_col + 1
    for col in range(2, total_col + 1):
        letter = get_column_letter(col)
        c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        c.font = TOTAL_FONT
        c.fill = TOTAL_FILL
    ws.cell(row=total_row, column=1).fill = TOTAL_FILL

    for wi in range(len(week_labels)):
        col_letter = get_column_letter(3 + wi)
        rng = f"{col_letter}{first_data_row}:{col_letter}{last_data_row}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"{col_letter}{first_data_row}>$B{first_data_row}"], fill=OVERLOAD_FILL))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND({col_letter}{first_data_row}>0,{col_letter}{first_data_row}<=$B{first_data_row})"],
            fill=OK_FILL))

    widths = [32, 13] + [15] * len(week_labels) + [13, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # charts
    chart_a = BarChart()
    chart_a.type = "col"
    chart_a.grouping = "clustered"
    chart_a.title = "Загрузка по неделям, по каждому исполнителю"
    chart_a.y_axis.title = "Часы"
    chart_a.height, chart_a.width = 10, 32
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    data = Reference(ws, min_col=3, max_col=last_week_col, min_row=header_row, max_row=last_data_row)
    chart_a.add_data(data, titles_from_data=True)
    chart_a.set_categories(cats)
    chart_a.gapWidth, chart_a.overlap = 60, -10
    ws.add_chart(chart_a, f"A{total_row + 3}")

    bar_b = BarChart()
    bar_b.type = "col"
    bar_b.title = "Суммарная загрузка команды по неделям vs норма"
    bar_b.y_axis.title = "Часы"
    bar_b.height, bar_b.width = 10, 16
    week_label_cats = Reference(ws, min_col=3, max_col=last_week_col, min_row=header_row, max_row=header_row)
    team_totals = Reference(ws, min_col=3, max_col=last_week_col, min_row=total_row, max_row=total_row)
    bar_b.series.append(Series(team_totals, title="Запланировано, ч"))
    bar_b.set_categories(week_label_cats)

    norm_col = total_col + 3
    for i in range(len(week_labels)):
        ws.cell(row=total_row, column=norm_col + i, value=f"=$B${total_row}")
    line_b = LineChart()
    norm_ref = Reference(ws, min_col=norm_col, max_col=norm_col + len(week_labels) - 1,
                          min_row=total_row, max_row=total_row)
    line_b.series.append(Series(norm_ref, title="Норма команды, ч"))
    line_b.y_axis.axId = 200
    bar_b += line_b
    ws.add_chart(bar_b, f"A{total_row + 23}")
    for i in range(len(week_labels)):
        ws.column_dimensions[get_column_letter(norm_col + i)].hidden = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.title("Еженедельное планирование загрузки команды")

uploaded = st.file_uploader("Загрузите файл с Agile-доски (.xlsx)", type="xlsx")

if not uploaded:
    st.info("Загрузите еженедельный экспорт с доски — приложение разберёт тикеты по исполнителям, "
            "вы впишете часы прямо здесь и увидите загрузку команды на графиках.")
    st.stop()

if "file_name" not in st.session_state or st.session_state.file_name != uploaded.name:
    people_data, week_labels = parse_workbook(uploaded.getvalue())
    st.session_state.file_name = uploaded.name
    st.session_state.people_data = {k: v.copy() for k, v in people_data.items()}
    st.session_state.week_labels = week_labels
    st.session_state.norms = {name: DEFAULT_NORM for name in people_data}

people_data = st.session_state.people_data
week_labels = st.session_state.week_labels

if not week_labels:
    st.error("Не удалось найти недельные колонки часов в файле. Проверьте структуру листов исполнителей.")
    st.stop()

tab_input, tab_dashboard = st.tabs(["Ввод часов", "Дашборд"])

with tab_input:
    st.caption("Часы по тикетам с доски впишите в нужную неделю. Строки «Доп. работа» — для задач, "
               "которых нет на доске: впишите название и часы. Новые строки можно добавлять кнопкой "
               "«+» внизу таблицы, лишние — удалять через иконку корзины слева от строки.")
    names = list(people_data.keys())
    selected = st.selectbox("Исполнитель", names)
    norm_val = st.number_input(f"Норма часов в неделю — {selected}",
                                min_value=0.0, value=float(st.session_state.norms[selected]), step=1.0)
    st.session_state.norms[selected] = norm_val

    df = people_data[selected]
    column_config = {
        "Тикет": st.column_config.TextColumn(),
        "Название": st.column_config.TextColumn(width="large"),
        "Колонка": st.column_config.TextColumn(),
        "Доска": st.column_config.TextColumn(),
        "Срок": st.column_config.DateColumn(format="DD.MM.YYYY"),
        "В работе": st.column_config.CheckboxColumn(),
    }
    for wl in week_labels:
        column_config[wl] = st.column_config.NumberColumn(clean_week_label(wl), min_value=0.0, step=0.5)

    edited = st.data_editor(df, column_config=column_config, hide_index=True,
                             use_container_width=True, num_rows="dynamic", key=f"editor_{selected}")

    # Newly added rows arrive with NaN/None in unfilled cells — normalise before aggregation/export.
    edited["Тикет"] = edited["Тикет"].fillna("")
    edited["Название"] = edited["Название"].fillna("")
    edited["Колонка"] = edited["Колонка"].fillna("Доп. работа")
    edited["Доска"] = edited["Доска"].fillna("")
    edited["В работе"] = edited["В работе"].fillna(False)
    for wl in week_labels:
        edited[wl] = edited[wl].fillna(0.0)

    people_data[selected] = edited
    st.session_state.people_data[selected] = edited

with tab_dashboard:
    rows = []
    for name, df in people_data.items():
        row = {"Исполнитель": name, "Норма, ч/нед": st.session_state.norms[name]}
        for wl in week_labels:
            row[clean_week_label(wl)] = df[wl].sum()
        rows.append(row)
    summary = pd.DataFrame(rows)
    week_cols = [clean_week_label(w) for w in week_labels]

    st.subheader("Сводка по команде")
    def highlight_overload(row):
        styles = []
        for col in summary.columns:
            if col in week_cols and row[col] > row["Норма, ч/нед"] > 0:
                styles.append("background-color:#F4CCCC")
            elif col in week_cols and 0 < row[col] <= row["Норма, ч/нед"]:
                styles.append("background-color:#D9EAD3")
            else:
                styles.append("")
        return styles

    st.dataframe(summary.style.apply(highlight_overload, axis=1), use_container_width=True, hide_index=True)

    col1, col2 = st.columns(2)
    with col1:
        long_df = summary.melt(id_vars=["Исполнитель"], value_vars=week_cols,
                                var_name="Неделя", value_name="Часы")
        fig = px.bar(long_df, x="Исполнитель", y="Часы", color="Неделя", barmode="group",
                     title="Загрузка по неделям, по каждому исполнителю")
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        team_total = summary[week_cols].sum()
        team_norm = summary["Норма, ч/нед"].sum()
        fig2 = go.Figure()
        fig2.add_bar(x=week_cols, y=team_total.values, name="Запланировано, ч")
        fig2.add_scatter(x=week_cols, y=[team_norm] * len(week_cols),
                          mode="lines+markers", name="Норма команды, ч")
        fig2.update_layout(title="Суммарная загрузка команды по неделям vs норма")
        st.plotly_chart(fig2, use_container_width=True)

    overloaded = summary[summary.apply(
        lambda r: any(r[c] > r["Норма, ч/нед"] > 0 for c in week_cols), axis=1)]
    if not overloaded.empty:
        st.warning("Перегружены на одной или нескольких неделях: " + ", ".join(overloaded["Исполнитель"]))
    else:
        st.success("Перегрузов не обнаружено.")

st.divider()
st.subheader("Экспорт")
st.caption("Скачайте заполненный файл в формате Excel — с формулами, сводкой и графиками, для архива.")
if st.button("Сформировать Excel"):
    data = build_export_workbook(people_data, week_labels, st.session_state.norms)
    st.download_button(
        "Скачать файл",
        data=data,
        file_name=f"Загрузка_команды_{datetime.now():%Y-%m-%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

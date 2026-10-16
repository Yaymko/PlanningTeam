"""Actual hours from the time-tracking export: parsing, matching people and comparing with the plan."""
import io
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core import FIXED_COLS  # noqa: E402
from fact import (  # noqa: E402
    build_comparison_workbook,
    compare_plan_fact,
    deviation_level,
    fact_breakdown,
    match_people,
    parse_fact_workbook,
    week_title,
)
from storage import Store  # noqa: E402

HEADER = ["Месяц", "Дата", "Подразделение", "Фамилия И.О.", "Сет. имя", "Направление деятельности",
          "Тема расчетов", "Проект", "Вид затрат", "Этап проекта", "Внутренний план работ",
          "Этап внутреннего плана", "Часы", "Часы по табелю", "Примечание"]
WEEKS = ["05.10–11.10 ч", "12.10–18.10 ч"]
TODAY = date(2026, 10, 16)


def _fact_row(day, name, project, kind, hours, note=""):
    return [datetime(2026, 10, 1), day, "Отдел", name, "LOGIN", "DIRECTUM", "Тема", project, kind,
            "Этап", "План работ", "Этап плана", hours, 0, note]


def _fact_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fact_df():
    return parse_fact_workbook(_fact_xlsx([
        _fact_row(datetime(2026, 10, 12), "Иванов Иван Петрович", "ОГВ Разработка", "Инвестиции", 6, "фича"),
        _fact_row(datetime(2026, 10, 13), "Иванов Иван Петрович", "Общефирменные", "Накладные", 2, "митинг"),
        _fact_row(datetime(2026, 10, 13), "Иванов Иван Петрович", "ОГВ Разработка", "Инвестиции", 0, "пусто"),
        _fact_row(datetime(2026, 10, 14), "Сидоров Олег Ильич", "Продажи", "Маркетинг", 5),
        _fact_row(datetime(2026, 10, 9), "Иванов Иван Петрович", "ОГВ Разработка", "Инвестиции", 3),
    ]))


def _people(hours_by_name):
    out = {}
    for name, (plan, fact) in hours_by_name.items():
        df = pd.DataFrame([{"Тикет": "ABC-1", "Название": "Задача", "Колонка": "10 Работа", "Доска": "1267",
                            "Срок": None, "В работе": True, WEEKS[0]: 0.0, WEEKS[1]: plan}],
                          columns=FIXED_COLS + WEEKS)
        df[f"{WEEKS[1][:-2]} факт, ч"] = fact
        out[name] = df
    return out


def test_parse_skips_zero_hours_and_keeps_groups():
    df = _fact_df()
    assert len(df) == 4
    assert df["Часы"].sum() == 16
    assert set(df["Проект"]) == {"ОГВ Разработка", "Общефирменные", "Продажи"}
    assert df.loc[0, "Дата"] == date(2026, 10, 12)


def test_match_people_by_surname_and_initials():
    plan = ["Иванов Иван", "Касьянов А.А.", "Петрова Анна Сергеевна"]
    got = match_people(["Иванов Иван Петрович", "Касьянов Андрей Александрович", "Петрова Анна Сергеевна",
                        "Сидоров Олег Ильич"], plan)
    assert got == {"Иванов Иван Петрович": "Иванов Иван", "Касьянов Андрей Александрович": "Касьянов А.А.",
                   "Петрова Анна Сергеевна": "Петрова Анна Сергеевна"}


def test_compare_plan_fact_by_person():
    store_people = {name: (df, 40.0, None) for name, df in _people({"Иванов Иван": (10.0, 7.0),
                                                                     "Петрова Анна": (8.0, 0.0)}).items()}
    week = _fact_df()
    week = week[week["Дата"] >= date(2026, 10, 12)]
    summary, fact = compare_plan_fact(store_people, WEEKS[1], week)
    rows = summary.set_index("Сотрудник")
    assert rows.loc["Иванов Иван", "План, ч"] == 10
    assert rows.loc["Иванов Иван", "Факт в планировщике, ч"] == 7
    assert rows.loc["Иванов Иван", "Факт по отметке, ч"] == 8
    assert rows.loc["Иванов Иван", "Отклонение, ч"] == -2
    assert rows.loc["Петрова Анна", "Факт по отметке, ч"] == 0
    assert pd.isna(rows.loc["Сидоров Олег Ильич", "Норма, ч"])   # only in the export
    assert set(fact["Сотрудник"]) == {"Иванов Иван", "Сидоров Олег Ильич"}

    breakdown = fact_breakdown(fact[fact["Сотрудник"] == "Иванов Иван"], "Проект")
    assert breakdown["Проект"].tolist() == ["ОГВ Разработка", "Общефирменные"]
    assert breakdown["Доля"].tolist() == [0.75, 0.25]
    assert breakdown["Примечание"].tolist() == ["фича", "митинг"]

    wb = openpyxl.load_workbook(io.BytesIO(build_comparison_workbook(summary, fact, "Проект")))
    assert wb.sheetnames == ["План vs факт", "Факт — Проект", "Отметки"]


def test_deviation_level():
    assert deviation_level(10, 10.5) == "ok"
    assert deviation_level(10, 8) == "warn"
    assert deviation_level(10, 5) == "bad"
    assert deviation_level(0, 3) == "bad"
    assert deviation_level(0, 0) == "ok"


def test_store_keeps_fact_by_week_and_finds_plan(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    store.save_plan("board.xlsx", _people({"Иванов Иван": (10.0, 0.0)}), WEEKS, today=TODAY)
    assert store.save_fact("fact.xlsx", _fact_df()) == ["2026-10-05", "2026-10-12"]
    assert [w["key"] for w in store.fact_weeks()] == ["2026-10-12", "2026-10-05"]
    assert week_title("2026-10-12") == "12.10–18.10"

    week = store.load_fact("2026-10-12")
    assert week["Часы"].sum() == 13
    assert week["Дата"].min() == date(2026, 10, 12)

    label, people = store.plan_for_week("2026-10-12", today=TODAY)
    assert label == WEEKS[1]
    assert list(people) == ["Иванов Иван"]
    assert store.plan_for_week("2026-09-28", today=TODAY) == (None, {})

    # A re-upload replaces the weeks it covers.
    store.save_fact("fact2.xlsx", _fact_df().iloc[:1])
    assert store.load_fact("2026-10-12")["Часы"].sum() == 6


def test_plan_for_week_from_history(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    store.save_plan("board.xlsx", _people({"Иванов Иван": (10.0, 0.0)}), WEEKS, today=TODAY)
    store.save_plan("board2.xlsx", _people({"Иванов Иван": (0.0, 0.0)}).copy(),
                    ["19.10–25.10 ч"], today=TODAY)
    label, people = store.plan_for_week("2026-10-12", today=TODAY)
    assert label == WEEKS[1]
    assert people["Иванов Иван"][0][label].sum() == 10

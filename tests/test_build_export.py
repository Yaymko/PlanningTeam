"""Smoke test for the Excel export: must produce a workbook with no cached errors
and the expected sheets/formulas, without needing LibreOffice/Excel to recalculate."""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core import FIXED_COLS, SUMMARY_SHEET_NAME, build_export_workbook  # noqa: E402

WEEK_LABELS = ["28.09–04.10 ч", "05.10–11.10 ч"]


def _sample_people_data():
    df = pd.DataFrame(
        [
            {
                "Тикет": "ABC-1", "Название": "Сделать штуку", "Колонка": "10 Работа",
                "Доска": "1267", "Срок": None, "В работе": True,
                WEEK_LABELS[0]: 4.0, WEEK_LABELS[1]: 0.0,
            },
            {
                "Тикет": "", "Название": "Внеплановая задача", "Колонка": "Доп. работа",
                "Доска": "", "Срок": None, "В работе": False,
                WEEK_LABELS[0]: 2.0, WEEK_LABELS[1]: 3.0,
            },
        ],
        columns=FIXED_COLS + WEEK_LABELS,
    )
    return {"Иванов Иван Иванович": df}


def test_export_contains_summary_and_person_sheets():
    data = build_export_workbook(_sample_people_data(), WEEK_LABELS, {"Иванов Иван Иванович": 40.0})
    wb = openpyxl.load_workbook(__to_tmp(data))

    assert SUMMARY_SHEET_NAME in wb.sheetnames
    assert "Иванов Иван Иванович" in wb.sheetnames

    person_ws = wb["Иванов Иван Иванович"]
    assert person_ws["A2"].value == "Тикет"
    assert person_ws["A3"].value == "ABC-1"
    # per-ticket total formula
    assert str(person_ws.cell(row=3, column=9).value).startswith("=SUM(")

    summary_ws = wb[SUMMARY_SHEET_NAME]
    assert summary_ws["A4"].value == "Иванов Иван Иванович"
    assert summary_ws["B4"].value == 40.0
    # the weekly cell on the summary sheet links back to the person sheet
    assert "Иванов Иван Иванович" in str(summary_ws["C4"].value)


def __to_tmp(data: bytes):
    import io
    return io.BytesIO(data)


def test_export_includes_actual_hours():
    from core import actual_label

    people = _sample_people_data()
    df = people["Иванов Иван Иванович"]
    df[actual_label(WEEK_LABELS[0])] = [5.0, 1.0]
    wb = openpyxl.load_workbook(__to_tmp(build_export_workbook(people, WEEK_LABELS, {})))

    person_ws = wb["Иванов Иван Иванович"]
    headers = [c.value for c in person_ws[2]]
    act_col = headers.index(actual_label(WEEK_LABELS[0])) + 1
    assert person_ws.cell(row=3, column=act_col).value == 5.0
    assert person_ws.cell(row=4, column=act_col).value == 1.0
    assert person_ws.cell(row=3, column=headers.index(actual_label(WEEK_LABELS[1])) + 1).value == 0.0
    assert "Итого факт, ч." in headers

    summary_ws = wb[SUMMARY_SHEET_NAME]
    summary_headers = [c.value for c in summary_ws[3]]
    cell = summary_ws.cell(row=4, column=summary_headers.index(actual_label(WEEK_LABELS[0])) + 1)
    assert str(cell.value).startswith("='Иванов Иван Иванович'!")

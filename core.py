"""
Чистая логика приложения — без зависимостей от Streamlit, легко тестируется.
Парсинг еженедельного экспорта с Agile-доски и сборка итогового Excel-файла.
"""
import io
import re
from datetime import date

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OVERVIEW_PREFIX = "Обзор тикетов"
SUMMARY_SHEET_NAME = "Сводка по команде"
FIXED_COLS = ["Тикет", "Название", "Колонка", "Доска", "Срок", "Статус", "Приоритет"]
DEFAULT_NORM = 40.0
ACTUAL_SUFFIX = " факт, ч"

STATUS_COL = "Статус"
PRIORITY_COL = "Приоритет"
LEGACY_IN_WORK_COL = "В работе"  # было до колонки «Статус»: галочка «в работе»
IN_WORK_STATUS = "в работе"
DEFAULT_STATUSES = [IN_WORK_STATUS, "пауза", "заморозка", "отложено"]
PRIORITIES = ["срочно", "высокий", "средний", "низкий"]  # от самого срочного


def status_from_legacy(in_work) -> str:
    """Галочка «В работе» из старых данных и выгрузок -> статус: отмечено — «в работе», иначе пусто."""
    if isinstance(in_work, str):
        return IN_WORK_STATUS if in_work.strip().lower() in ("да", "true", "1", "+", "x", IN_WORK_STATUS) else ""
    try:
        return IN_WORK_STATUS if in_work and not pd.isna(in_work) else ""
    except (TypeError, ValueError):
        return IN_WORK_STATUS if in_work else ""


def sort_by_priority(df: pd.DataFrame) -> pd.DataFrame:
    """Строки по приоритету: «срочно» наверху, без приоритета — внизу; внутри одного приоритета
    порядок не меняется."""
    if PRIORITY_COL not in df.columns or df.empty:
        return df
    rank = df[PRIORITY_COL].map(lambda p: PRIORITIES.index(p) if p in PRIORITIES else len(PRIORITIES))
    return df.iloc[rank.argsort(kind="stable")].reset_index(drop=True)


def actual_label(week_label: str) -> str:
    """Колонка фактических часов для недели: «28.09–04.10 ч» -> «28.09–04.10 факт, ч»."""
    return re.sub(r"\s*ч\.?$", "", str(week_label)).strip() + ACTUAL_SUFFIX


def is_actual_label(col) -> bool:
    return str(col).endswith(ACTUAL_SUFFIX)


def hour_cols(week_labels: list) -> list:
    """Колонки часов по неделям: план и рядом факт для каждой недели."""
    cols = []
    for wl in week_labels:
        cols += [wl, actual_label(wl)]
    return cols


def parse_workbook(file_bytes: bytes):
    """Разбирает еженедельный экспорт с доски: лист на исполнителя -> DataFrame тикетов."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    person_sheets = [
        s for s in wb.sheetnames
        if not s.startswith(OVERVIEW_PREFIX) and s != SUMMARY_SHEET_NAME
    ]

    people_data = {}
    week_labels = None

    for name in person_sheets:
        ws = wb[name]
        # Заголовок обычно во 2-й строке (над ним название листа), но в некоторых выгрузках
        # строки с названием нет и заголовок стоит первым — ищем его в первых строках.
        header_row = next(
            (r for r in range(1, min(ws.max_row, 5) + 1) if ws.cell(row=r, column=1).value == "Тикет"),
            None,
        )
        if header_row is None:
            continue
        header = [c.value for c in ws[header_row]]
        idx = {h: i for i, h in enumerate(header) if h is not None}
        # Недели идут сразу после последней из известных колонок тикета («В работе» в старых
        # выгрузках, «Статус»/«Приоритет» — если их добавили в файл).
        known = [idx[h] for h in [*FIXED_COLS, LEGACY_IN_WORK_COL] if h in idx]
        start_week_col = max(known) + 1
        end_week_col = idx.get("Итого, ч.", len(header) - 1)
        this_week_labels = [header[i] for i in range(start_week_col, end_week_col)]
        if week_labels is None:
            week_labels = this_week_labels

        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            a_val = ws.cell(row=r, column=1).value
            c_val = ws.cell(row=r, column=idx["Колонка"] + 1).value
            if isinstance(a_val, str) and a_val.strip().startswith("Итого"):  # «Итого по неделям:» или «Итого»
                break
            row = {
                "Тикет": a_val or "",
                "Название": ws.cell(row=r, column=idx["Название"] + 1).value or "",
                "Колонка": c_val or "",
                "Доска": ws.cell(row=r, column=idx.get("Доска", 2) + 1).value if "Доска" in idx else "",
                "Срок": ws.cell(row=r, column=idx["Срок"] + 1).value,
            }
            if STATUS_COL in idx:
                row[STATUS_COL] = ws.cell(row=r, column=idx[STATUS_COL] + 1).value or ""
            elif LEGACY_IN_WORK_COL in idx:
                row[STATUS_COL] = status_from_legacy(ws.cell(row=r, column=idx[LEGACY_IN_WORK_COL] + 1).value)
            if PRIORITY_COL in idx:
                row[PRIORITY_COL] = ws.cell(row=r, column=idx[PRIORITY_COL] + 1).value or ""
            for wi, wl in enumerate(this_week_labels):
                v = ws.cell(row=r, column=start_week_col + 1 + wi).value
                row[wl] = float(v) if isinstance(v, (int, float)) else 0.0
            rows.append(row)

        df = pd.DataFrame(rows, columns=FIXED_COLS + this_week_labels)
        people_data[name] = df

    return people_data, (week_labels or [])


def normalize_rows(df: pd.DataFrame, week_labels: list) -> pd.DataFrame:
    """Приводит таблицу тикетов к единому виду: пустые ячейки новых строк -> значения по умолчанию,
    часы -> float. Нужна после data_editor (добавленные строки приходят с NaN/None).
    Фактических часов может не быть (выгрузка с доски, данные до их появления) — тогда 0."""
    df = df.copy()
    if STATUS_COL not in df.columns and LEGACY_IN_WORK_COL in df.columns:
        df[STATUS_COL] = df[LEGACY_IN_WORK_COL].map(status_from_legacy)
    df = df.reindex(columns=FIXED_COLS + hour_cols(week_labels))
    df["Тикет"] = df["Тикет"].fillna("").astype(str)
    df["Название"] = df["Название"].fillna("").astype(str)
    df["Колонка"] = df["Колонка"].fillna("Доп. работа").astype(str)
    df["Доска"] = df["Доска"].fillna("").astype(str)
    df["Срок"] = df["Срок"].astype(object).where(df["Срок"].notna(), None)
    for col in (STATUS_COL, PRIORITY_COL):
        df[col] = df[col].map(lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip())
    for col in hour_cols(week_labels):
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0).astype(float)
    return df.reset_index(drop=True)


DELETE_COL = "Удалить"


def drop_marked_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Убирает строки, отмеченные галочкой в колонке «Удалить», и саму колонку."""
    if DELETE_COL not in df.columns:
        return df
    marked = df[DELETE_COL].fillna(False).astype(bool)
    return df[~marked].drop(columns=DELETE_COL).reset_index(drop=True)


def drop_empty_rows(df: pd.DataFrame, week_labels: list) -> pd.DataFrame:
    """Убирает пустые строки: без тикета, без названия и без часов (план и факт) по всем неделям.
    Такие строки-заготовки «Доп. работа» приходят в каждой выгрузке с доски; новые строки
    сотрудник добавляет кнопкой «+» в таблице."""
    df = normalize_rows(df, week_labels)
    hours = df[hour_cols(week_labels)].abs().sum(axis=1) if week_labels else 0
    empty = (df["Тикет"].str.strip() == "") & (df["Название"].str.strip() == "") & (hours == 0)
    return df[~empty].reset_index(drop=True)


def merge_saved_hours(new_df: pd.DataFrame, old_df: pd.DataFrame, week_labels: list) -> pd.DataFrame:
    """Переносит уже сохранённые часы в свежую выгрузку с доски.

    Часы (план и факт) по тикетам переносятся по номеру тикета для недель, которые есть в обеих
    выгрузках, статус и приоритет — если были заданы. Строки «Доп. работа» с названием, которых нет
    в новой выгрузке, добавляются в конец.
    """
    new_df = normalize_rows(new_df, week_labels)
    common_weeks = [wl for wl in week_labels if wl in old_df.columns]
    if old_df.empty:
        return new_df
    not_week = {*FIXED_COLS, LEGACY_IN_WORK_COL}
    old_df = normalize_rows(old_df, [c for c in old_df.columns if c not in not_week and not is_actual_label(c)])

    old_by_ticket = {}
    for _, row in old_df.iterrows():
        if row["Тикет"]:
            old_by_ticket.setdefault(row["Тикет"], row)
    for i, ticket in new_df["Тикет"].items():
        if ticket in old_by_ticket:
            old = old_by_ticket[ticket]
            # Статус и приоритет ставят в планировщике — они важнее отметки из выгрузки.
            for col in (STATUS_COL, PRIORITY_COL):
                if old[col]:
                    new_df.at[i, col] = old[col]
            for col in hour_cols(common_weeks):
                new_df.at[i, col] = float(old[col])
    if not common_weeks:
        return new_df

    existing_titles = set(new_df.loc[new_df["Тикет"] == "", "Название"])
    extra = old_df[(old_df["Тикет"] == "") & (old_df["Название"] != "")
                   & ~old_df["Название"].isin(existing_titles)]
    if not extra.empty:
        extra = normalize_rows(extra, week_labels)
        new_df = pd.concat([new_df, extra], ignore_index=True)
    return new_df


def clean_week_label(label: str) -> str:
    return re.sub(r"\s*ч\.?$", "", str(label)).strip()


def _week_range(label: str, today: date):
    """(начало, конец) недели из заголовка вида «28.09–04.10 ч». Года в заголовке нет:
    берётся тот, при котором неделя ближе всего к сегодняшней дате."""
    m = re.match(r"\s*(\d{1,2})\.(\d{1,2})\s*[–—-]\s*(\d{1,2})\.(\d{1,2})", str(label))
    if not m:
        return None
    d1, m1, d2, m2 = map(int, m.groups())
    best = None
    for year in (today.year - 1, today.year, today.year + 1):
        try:
            start = date(year, m1, d1)
            end = date(year + (1 if m2 < m1 else 0), m2, d2)
        except ValueError:
            continue
        if best is None or abs((start - today).days) < abs((best[0] - today).days):
            best = (start, end)
    return best


def week_key(label: str, today: date) -> str:
    """Ключ недели для истории: дата начала (YYYY-MM-DD), чтобы недели разных лет не путались.
    Если дату из заголовка не прочитать — сам заголовок."""
    rng = _week_range(label, today)
    return rng[0].isoformat() if rng else clean_week_label(label)


def nearest_weeks(week_labels: list, today: date, count: int = 3) -> list:
    """Ближайшие недели плана: текущая и следующие за ней (всего count).

    Если до конца плана недель меньше count — последние count; если сегодня раньше всех недель —
    первые count, если позже всех — последние count. Если даты из заголовков не читаются — все недели.
    """
    ranges = [_week_range(wl, today) for wl in week_labels]
    if not week_labels or any(r is None for r in ranges):
        return list(week_labels)
    current = next((i for i, (_, end) in enumerate(ranges) if end >= today), len(week_labels))
    start = max(0, min(current, len(week_labels) - count))
    return list(week_labels[start:start + count])


def build_export_workbook(people_data: dict, week_labels: list, norms: dict) -> bytes:
    """Собирает итоговый Excel: лист на исполнителя + сводка по команде с формулами и графиками."""
    HEADER_FILL = PatternFill("solid", fgColor="2F4A6D")
    HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
    INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
    INPUT_FONT = Font(name="Arial", color="0000FF")
    LINK_FONT = Font(name="Arial", color="008000")
    TOTAL_FILL = PatternFill("solid", fgColor="DDE5EE")
    TOTAL_FONT = Font(name="Arial", bold=True)
    BODY_FONT = Font(name="Arial")
    TITLE_FONT = Font(name="Arial", bold=True, size=14)
    OVERLOAD_FILL = PatternFill(bgColor="F4CCCC")
    OK_FILL = PatternFill(bgColor="D9EAD3")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    week_start_col = len(FIXED_COLS) + 1
    total_col = week_start_col + len(week_labels)          # «Итого, ч.» по плану
    act_start_col = total_col + 1                           # факт по неделям
    act_total_col = act_start_col + len(week_labels)        # «Итого факт, ч.»
    actual_labels = [actual_label(wl) for wl in week_labels]

    totals_row = {}
    for name, df in people_data.items():
        df = sort_by_priority(normalize_rows(df, week_labels))
        ws = wb.create_sheet(name[:31])
        ws.cell(row=1, column=1, value=f"{name} — часы по неделям (план и факт), заполнено в планировщике")
        headers = FIXED_COLS + week_labels + ["Итого, ч."] + actual_labels + ["Итого факт, ч."]
        for ci, h in enumerate(headers, start=1):
            c = ws.cell(row=2, column=ci, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

        r = 3
        for _, row in df.iterrows():
            for ci, col in enumerate(FIXED_COLS, start=1):
                ws.cell(row=r, column=ci, value=row[col])
            for wi, wl in enumerate(week_labels):
                ws.cell(row=r, column=week_start_col + wi, value=row[wl])
                ws.cell(row=r, column=act_start_col + wi, value=row[actual_label(wl)])
            for first, last in ((week_start_col, total_col), (act_start_col, act_total_col)):
                ws.cell(row=r, column=last,
                        value=f"=SUM({get_column_letter(first)}{r}:{get_column_letter(last - 1)}{r})")
            r += 1

        last_data_row = r - 1
        ws.cell(row=r, column=1, value="Итого по неделям:").font = TOTAL_FONT
        for ci in range(week_start_col, act_total_col + 1):
            letter = get_column_letter(ci)
            c = ws.cell(row=r, column=ci, value=f"=SUM({letter}3:{letter}{last_data_row})")
            c.font = TOTAL_FONT
            c.fill = TOTAL_FILL
        ws.cell(row=r, column=1).fill = TOTAL_FILL
        totals_row[name] = r

        widths = [14, 50, 22, 10, 12, 12, 11] + ([14] * len(week_labels) + [12]) * 2
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --- summary sheet ---
    ws = wb.create_sheet(SUMMARY_SHEET_NAME, 0)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "Еженедельное планирование загрузки команды"
    ws["A1"].font = TITLE_FONT

    header_row = 3
    headers = (["Исполнитель", "Норма, ч/нед"] + week_labels + ["Итого, ч", "Загрузка посл. недели"]
               + actual_labels + ["Итого факт, ч"])
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    last_week_col = 2 + len(week_labels)
    sum_total_col = last_week_col + 1
    load_col = sum_total_col + 1
    sum_act_col = load_col + 1
    sum_act_total_col = sum_act_col + len(week_labels)

    first_data_row = header_row + 1
    names = list(people_data.keys())
    for idx, name in enumerate(names):
        r = first_data_row + idx
        ws.cell(row=r, column=1, value=name).font = BODY_FONT
        norm_cell = ws.cell(row=r, column=2, value=norms.get(name, DEFAULT_NORM))
        norm_cell.fill = INPUT_FILL
        norm_cell.font = INPUT_FONT
        tr = totals_row[name]
        for wi in range(len(week_labels)):
            for dst, src in ((3 + wi, week_start_col + wi), (sum_act_col + wi, act_start_col + wi)):
                cell = ws.cell(row=r, column=dst)
                cell.value = f"='{name}'!{get_column_letter(src)}{tr}"
                cell.font = LINK_FONT
                cell.number_format = "0.##"
        ws.cell(row=r, column=sum_total_col,
                value=f"=SUM(C{r}:{get_column_letter(last_week_col)}{r})").font = BODY_FONT
        ws.cell(row=r, column=load_col,
                value=f'=IF(B{r}=0,"",{get_column_letter(last_week_col)}{r}/B{r})').number_format = "0%"
        ws.cell(row=r, column=sum_act_total_col,
                value=f"=SUM({get_column_letter(sum_act_col)}{r}:"
                      f"{get_column_letter(sum_act_total_col - 1)}{r})").font = BODY_FONT

    last_data_row = first_data_row + len(names) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="Итого по команде:").font = TOTAL_FONT
    for col in [*range(2, sum_total_col + 1), *range(sum_act_col, sum_act_total_col + 1)]:
        letter = get_column_letter(col)
        c = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        c.font = TOTAL_FONT
        c.fill = TOTAL_FILL
    ws.cell(row=total_row, column=1).fill = TOTAL_FILL

    for wi in range(len(week_labels)):
        col_letter = get_column_letter(3 + wi)
        rng = f"{col_letter}{first_data_row}:{col_letter}{last_data_row}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"{col_letter}{first_data_row}>$B{first_data_row}"], fill=OVERLOAD_FILL))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND({col_letter}{first_data_row}>0,{col_letter}{first_data_row}<=$B{first_data_row})"],
            fill=OK_FILL))

    widths = [32, 13] + [15] * len(week_labels) + [13, 18] + [15] * len(week_labels) + [13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # charts
    chart_a = BarChart()
    chart_a.type = "col"
    chart_a.grouping = "clustered"
    chart_a.title = "Загрузка по неделям, по каждому исполнителю"
    chart_a.y_axis.title = "Часы"
    chart_a.height, chart_a.width = 10, 32
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    data = Reference(ws, min_col=3, max_col=last_week_col, min_row=header_row, max_row=last_data_row)
    chart_a.add_data(data, titles_from_data=True)
    chart_a.set_categories(cats)
    chart_a.gapWidth, chart_a.overlap = 60, -10
    ws.add_chart(chart_a, f"A{total_row + 3}")

    bar_b = BarChart()
    bar_b.type = "col"
    bar_b.title = "Суммарная загрузка команды по неделям vs норма"
    bar_b.y_axis.title = "Часы"
    bar_b.height, bar_b.width = 10, 16
    week_label_cats = Reference(ws, min_col=3, max_col=last_week_col, min_row=header_row, max_row=header_row)
    team_totals = Reference(ws, min_col=3, max_col=last_week_col, min_row=total_row, max_row=total_row)
    bar_b.series.append(Series(team_totals, title="Запланировано, ч"))
    team_actuals = Reference(ws, min_col=sum_act_col, max_col=sum_act_total_col - 1,
                             min_row=total_row, max_row=total_row)
    bar_b.series.append(Series(team_actuals, title="Факт, ч"))
    bar_b.set_categories(week_label_cats)

    norm_col = sum_act_total_col + 2
    for i in range(len(week_labels)):
        ws.cell(row=total_row, column=norm_col + i, value=f"=$B${total_row}")
    line_b = LineChart()
    norm_ref = Reference(ws, min_col=norm_col, max_col=norm_col + len(week_labels) - 1,
                          min_row=total_row, max_row=total_row)
    line_b.series.append(Series(norm_ref, title="Норма команды, ч"))
    line_b.y_axis.axId = 200
    bar_b += line_b
    ws.add_chart(bar_b, f"A{total_row + 23}")
    for i in range(len(week_labels)):
        ws.column_dimensions[get_column_letter(norm_col + i)].hidden = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
